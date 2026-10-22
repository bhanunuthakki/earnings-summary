"""MELI workbook formula specification and verified artifact-engine export.

The numerical oracle remains meli_model. This owner defines spreadsheet formulas
once; the replaceable JS adapter contains file and calculation mechanics only.
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import shutil
import subprocess
import tempfile
from dataclasses import asdict, dataclass, replace
from pathlib import Path
from typing import Literal, TypedDict

from openpyxl.utils import get_column_letter

from dcf.input_evidence import InputEvidenceError, ModelInputReceipt, canonical_digest
from dcf.meli_inputs import (
    BRIDGE_SENSITIVITY_CASES,
    FINANCING_RECIPE,
    MELI_DIAGNOSTIC_UNITS,
    RECIPE,
    effective_numeric_inputs,
    is_meli_recipe,
    model_output,
)
from dcf.meli_model import (
    Assum,
    Mirror,
    implied_credit_terminal_roe,
    mirror,
    validate_credit_terminal,
)
from dcf.readiness import replay_model_scenarios
from dcf.redesign import GROWTH_FADE_CURVATURE as _CURVATURE

PCT = "0.0%"
USD0 = "#,##0"
NUM1 = "0.0"
NUM2 = "0.00"
MULT = '0.0"x"'

# Dashboard addresses remain stable for existing readers. Source anchors are
# fixed on Actuals; forecast drivers select one case and feed one Model build.
R = {
    "comm_rev0": 3,
    "comm_gn": 4,
    "comm_gt": 5,
    "comm_mn": 6,
    "comm_mt": 7,
    "fpay_rev0": 8,
    "fpay_gn": 9,
    "fpay_gt": 10,
    "fpay_mn": 11,
    "fpay_mt": 12,
    "da": 13,
    "capxn": 14,
    "capxt": 15,
    "nwc": 16,
    "tax": 17,
    "opmult": 18,
    "wacc": 19,
    "cb0": 20,
    "cbgn": 21,
    "cbgt": 22,
    "niman": 23,
    "nimat": 24,
    "cropex": 25,
    "cap": 26,
    "cke": 27,
    "cgt": 28,
    "croe": 29,
    "netcash": 30,
    "years": 31,
    "sh": 32,
    "px": 33,
    "curv": 34,
    "opev": 37,
    "crev": 38,
    "eqv": 39,
    "vps": 40,
    "up": 41,
    "termrev": 42,
    "croet": 43,
}
WORKBOOK_DRIVERS = (
    ("comm_rev0", "comm_rev0", "Commerce revenue baseline ($M)", USD0),
    ("comm_gn", "comm_g_near", "Commerce growth near", PCT),
    ("comm_gt", "comm_g_term", "Commerce growth terminal", PCT),
    ("comm_mn", "comm_margin_near", "Commerce EBIT margin near", PCT),
    ("comm_mt", "comm_margin_term", "Commerce EBIT margin terminal", PCT),
    ("fpay_rev0", "fpay_rev0", "Payments revenue baseline ($M; calculated)", USD0),
    ("fpay_gn", "fpay_g_near", "Payments growth near", PCT),
    ("fpay_gt", "fpay_g_term", "Payments growth terminal", PCT),
    ("fpay_mn", "fpay_margin_near", "Payments EBIT margin near", PCT),
    ("fpay_mt", "fpay_margin_term", "Payments EBIT margin terminal", PCT),
    ("da", "da_pct", "D&A / operating revenue", PCT),
    ("capxn", "capex_pct_near", "Capex / revenue near", PCT),
    ("capxt", "capex_pct_term", "Capex / revenue terminal", PCT),
    ("nwc", "nwc_pct", "Working capital / incremental revenue", PCT),
    ("tax", "tax", "Tax rate", PCT),
    ("opmult", "op_exit_ebitda_mult", "Operating exit EV / EBITDA", MULT),
    ("wacc", "wacc", "Operating discount rate", PCT),
    ("cb0", "cb0", "Gross credit book baseline ($M)", USD0),
    ("cbgn", "cbg_near", "Credit book growth near", PCT),
    ("cbgt", "cbg_term", "Credit book growth terminal", PCT),
    ("niman", "nimal_near", "NIMAL near", PCT),
    ("nimat", "nimal_term", "NIMAL terminal", PCT),
    ("cropex", "credit_opex_ratio", "Credit operating costs / book", PCT),
    ("cap", "cap_ratio", "Required equity / credit book", PCT),
    ("cke", "credit_ke", "Credit cost of equity", PCT),
    ("cgt", "credit_g_term", "Credit terminal growth", PCT),
    ("croe", "credit_terminal_roe", "Credit terminal ROE", PCT),
    ("netcash", "net_cash", "Net non-operating cash ($M; calculated)", USD0),
    ("years", "years", "Forecast years (fixed at build)", "0"),
    ("sh", "shares", "Valuation shares assumption (M)", NUM2),
    ("px", "price", "Quote / share ($; unavailable if zero)", NUM2),
)


class WorkbookSheetSpec(TypedDict):
    name: str
    cells: dict[str, str | float | int]
    formats: dict[str, str]
    inputs: list[str]


BearProvenance = Literal["thesis", "seed", "explicit_vector_acceptance_separate"]
BEAR_PROVENANCE_NOTES: dict[BearProvenance, str] = {
    "thesis": "thesis (holdings bear_deltas)",
    "seed": "seed fallback (generic scenario deltas)",
    "explicit_vector_acceptance_separate": "Explicit vectors; acceptance recorded separately",
}


class MeliWorkbookSpec(TypedDict):
    schema: str
    sheets: list[WorkbookSheetSpec]
    data_tables: list[dict[str, str]]
    expected: dict[str, float]
    input_receipt_sha256: str | None
    scenario_commitment_sha256: str
    input_recipe: str
    bear_provenance_note: str


@dataclass(frozen=True)
class WorkbookRuntime:
    """Explicit loader-supplied runtime paths; no product runtime discovery."""

    node: Path
    node_modules: Path

    @classmethod
    def from_environment(cls) -> WorkbookRuntime:
        node = os.environ.get("DCF_ARTIFACT_NODE", "")
        modules = os.environ.get("DCF_ARTIFACT_NODE_MODULES", "")
        if not node or not modules:
            raise InputEvidenceError("workbook_artifact_runtime_not_configured")
        runtime = cls(Path(node), Path(modules))
        if not runtime.node.is_file() or not runtime.node_modules.is_dir():
            raise InputEvidenceError("workbook_artifact_runtime_unavailable")
        return runtime


def validate_cases(
    s: Assum,
    cases: dict[str, Assum],
    *,
    recipe: str = RECIPE,
) -> dict[str, Assum]:
    if not is_meli_recipe(recipe):
        raise InputEvidenceError("model_input_recipe_mismatch")
    if set(cases) != {"base", "bear", "bull"} or not 1 <= s.years <= 100:
        raise InputEvidenceError("workbook_case_population_or_horizon_invalid")
    base_values = asdict(s)
    if any(getattr(cases["base"], attr) != base_values[attr] for _, attr, _, _ in WORKBOOK_DRIVERS):
        raise InputEvidenceError("workbook_base_case_input_mismatch")
    for case in cases.values():
        values = asdict(case)
        if any(not math.isfinite(float(values[attr])) for _, attr, _, _ in WORKBOOK_DRIVERS):
            raise InputEvidenceError("workbook_nonfinite_input")
        validate_credit_terminal(values)
        fixed = ("comm_rev0", "fpay_rev0", "cb0", "shares", "net_cash")
        if recipe == FINANCING_RECIPE:
            fixed += ("financing_other_expenses_reserve",)
        if case.years != s.years or any(getattr(case, key) != getattr(s, key) for key in fixed):
            raise InputEvidenceError("workbook_case_reported_anchor_changed")
    return cases


def workbook_spec(
    s: Assum,
    m: Mirror,
    *,
    cases: dict[str, Assum],
    input_receipt: ModelInputReceipt | None = None,
    ticker: str = "MELI",
    bear_provenance: BearProvenance = "explicit_vector_acceptance_separate",
) -> MeliWorkbookSpec:
    """One formula owner; the JS adapter supplies calculation and file mechanics."""
    recipe = input_receipt.recipe if input_receipt is not None else RECIPE
    cases = validate_cases(s, cases, recipe=recipe)
    if not math.isclose(m.vps, mirror(s).vps, rel_tol=1e-12, abs_tol=1e-9):
        raise InputEvidenceError("model_output_input_mismatch")
    sheets: dict[str, WorkbookSheetSpec] = {
        name: {"name": name, "cells": {}, "formats": {}, "inputs": []}
        for name in ("Dashboard", "Assumptions", "Model", "Valuation", "Scenario", "Actuals")
    }

    def put(name: str, address: str, value: str | float | int, fmt: str = "") -> None:
        sheets[name]["cells"][address] = value
        if fmt:
            sheets[name]["formats"][address] = fmt

    put("Dashboard", "A1", f"{ticker} platform DCF")
    put("Dashboard", "D2", "Case selected:")
    put("Dashboard", "E2", '=CHOOSE(B44,"Base","Bear","Bull")')
    put("Dashboard", "A2", "Operating FCFF plus credit FCFE ($M)")
    put("Assumptions", "A1", "Forecast assumptions by case")
    put("Assumptions", "B2", "Base")
    put("Assumptions", "C2", "Bear")
    put("Assumptions", "D2", "Bull")
    put("Actuals", "A1", "Fixed financial baseline: reported inputs and stated calculations")
    for index, name in enumerate(("base", "bear", "bull"), 2):
        column = get_column_letter(index)
        for key, attr, label, fmt in WORKBOOK_DRIVERS:
            row = R[key]
            put("Assumptions", f"A{row}", label)
            put("Dashboard", f"A{row}", label)
            if key in {"comm_rev0", "fpay_rev0", "cb0", "netcash", "years"}:
                put("Actuals", f"A{row}", label)
                put("Actuals", f"B{row}", getattr(s, attr), fmt)
                put("Dashboard", f"B{row}", f"='Actuals'!B{row}", fmt)
                put("Assumptions", f"{column}{row}", f"='Actuals'!B{row}", fmt)
            elif key == "sh":
                put("Actuals", "A32", "Reported diluted shares baseline (M)")
                put("Actuals", "B32", s.shares, NUM2)
                put("Dashboard", "B32", s.shares, NUM2)
                put("Assumptions", f"{column}{row}", "='Actuals'!B32", fmt)
            else:
                put("Assumptions", f"{column}{row}", getattr(cases[name], attr), fmt)
                sheets["Assumptions"]["inputs"].append(f"{column}{row}")
                selected = (
                    f"CHOOSE($B$44,'Assumptions'!B{row},'Assumptions'!C{row},'Assumptions'!D{row})"
                )
                if key in {"wacc", "cke"}:
                    selected = f'IF(D{row}="",{selected},D{row})'
                put("Dashboard", f"B{row}", "=" + selected, fmt)
    put("Dashboard", "A34", "Growth fade curvature")
    put("Dashboard", "B34", _CURVATURE, NUM1)
    put("Dashboard", "A44", "Case index (1 Base; 2 Bear; 3 Bull)")
    put("Dashboard", "B44", 1, "0")
    put("Dashboard", "D18", "Optional discount-rate overrides")
    put("Dashboard", "C19", "WACC override")
    put("Dashboard", "C27", "Credit Ke override")
    sheets["Dashboard"]["inputs"] = ["B32", "B44", "D19", "D27"]
    sheets["Dashboard"]["formats"].update({"D19": PCT, "D27": PCT})
    put("Actuals", "A36", "Source binding")
    digest = canonical_digest(input_receipt.model_dump(mode="json")) if input_receipt else None
    put("Actuals", "B36", digest or "Unverified inputs; no financial acceptance")
    put("Actuals", "A37", "Research snapshot")
    put(
        "Actuals",
        "B37",
        input_receipt.request.research_snapshot_id if input_receipt else "Unavailable",
    )
    scenario_commitment = canonical_digest(
        replay_model_scenarios(
            "MELI",
            effective_numeric_inputs(asdict(s), recipe=recipe),
            model_output(effective_numeric_inputs(asdict(s), recipe=recipe)),
            {
                "reviewed_scenario_inputs": {
                    name: effective_numeric_inputs(asdict(case), recipe=recipe)
                    for name, case in cases.items()
                }
            },
            recipe=recipe,
        )
    )
    put("Actuals", "A38", "Scenario output commitment")
    put("Actuals", "B38", scenario_commitment)
    put("Actuals", "A39", "Financial period end")
    put(
        "Actuals",
        "B39",
        input_receipt.request.financial_period_end.isoformat() if input_receipt else "Unavailable",
    )
    put("Actuals", "A40", "Input recipe")
    put("Actuals", "B40", input_receipt.recipe if input_receipt else "Unverified inputs")
    if recipe == FINANCING_RECIPE:
        put("Actuals", "A41", "Financing expense reserve ($M; analyst input)")
        put("Actuals", "B41", s.financing_other_expenses_reserve, NUM2)
    for name, label, value in (
        ("Model", "B1", "C1"),
        ("Valuation", "A1", "B1"),
        ("Scenario", "A7", "B7"),
        ("Assumptions", "F1", "G1"),
    ):
        put(name, label, "Case selected:")
        put(name, value, "='Dashboard'!E2")
    n = s.years
    labels = {
        3: "Year",
        4: "Commerce revenue",
        5: "Payments revenue",
        6: "Operating revenue",
        7: "Operating EBIT",
        8: "D&A",
        9: "Capex",
        10: "Operating FCFF",
        11: "WACC factor",
        12: "PV operating FCFF",
        13: "Credit book",
        14: "Credit net income",
        15: "Required equity",
        16: "Credit FCFE",
        17: "Credit Ke factor",
        18: "PV credit FCFE",
    }
    for row, label in labels.items():
        put("Model", f"B{row}", label)

    def dref(key: str) -> str:
        return f"'Dashboard'!$B${R[key]}"

    def progression(c: str, near: str, term: str, *, fade: bool = False) -> str:
        a, b, years = dref(near), dref(term), dref("years")
        if fade:
            return f"IF({years}<=1,{a},{b}+({a}-{b})*(({years}-{c}$3)/({years}-1))^{dref('curv')})"
        return f"IF({years}<=1,{a},{a}+({b}-{a})*({c}$3-1)/({years}-1))"

    for year in range(1, n + 1):
        c, p = get_column_letter(year + 2), get_column_letter(year + 1)
        put("Model", f"{c}3", year)
        prev_op = f"({dref('comm_rev0')}+{dref('fpay_rev0')})" if year == 1 else f"{p}6"
        prev_cb = dref("cb0") if year == 1 else f"{p}13"
        prev_cap = f"{dref('cap')}*{dref('cb0')}" if year == 1 else f"{p}15"
        for row, near, term, anchor in (
            (4, "comm_gn", "comm_gt", "comm_rev0"),
            (5, "fpay_gn", "fpay_gt", "fpay_rev0"),
            (13, "cbgn", "cbgt", "cb0"),
        ):
            previous = dref(anchor) if year == 1 else f"{p}{row}"
            put(
                "Model",
                f"{c}{row}",
                f"={previous}*(1+{progression(c, near, term, fade=True)})",
                USD0,
            )
        formulas = {
            6: f"{c}4+{c}5",
            7: f"{c}4*({progression(c, 'comm_mn', 'comm_mt')})+{c}5*({progression(c, 'fpay_mn', 'fpay_mt')})",
            8: f"{c}6*{dref('da')}",
            9: f"{c}6*({progression(c, 'capxn', 'capxt')})",
            10: f"{c}7*(1-{dref('tax')})+{c}8-{c}9-({c}6-{prev_op})*{dref('nwc')}",
            11: f"1/(1+{dref('wacc')})^{c}3",
            12: f"{c}10*{c}11",
            14: f"(({c}13+{prev_cb})/2)*(({progression(c, 'niman', 'nimat')})-{dref('cropex')})*(1-{dref('tax')})",
            15: f"{dref('cap')}*{c}13",
            16: f"{c}14-({c}15-{prev_cap})",
            17: f"1/(1+{dref('cke')})^{c}3",
            18: f"{c}16*{c}17",
        }
        for row, formula in formulas.items():
            put("Model", f"{c}{row}", "=" + formula, NUM2 if row in {11, 17} else USD0)
    end = get_column_letter(n + 2)
    valuation = {
        2: ("PV operating FCFF", f"SUM('Model'!C12:{end}12)"),
        3: ("Operating terminal EV", f"('Model'!{end}7+'Model'!{end}8)*{dref('opmult')}"),
        4: ("PV operating terminal", f"B3*'Model'!{end}11"),
        5: ("Operating EV", "B2+B4"),
        6: ("PV credit FCFE", f"SUM('Model'!C18:{end}18)"),
        7: (
            "Credit terminal value",
            f"'Model'!{end}14*(1+{dref('cgt')})*(1-{dref('cgt')}/{dref('croe')})/({dref('cke')}-{dref('cgt')})",
        ),
        8: ("PV credit terminal", f"B7*'Model'!{end}17"),
        9: ("Credit equity value", "B6+B8"),
        10: ("Net non-operating cash", dref("netcash")),
        11: ("Equity value", "B5+B9+B10"),
        12: ("Value / valuation share ($)", f'IF(B16="Valid",B11/{dref("sh")},NA())'),
        13: ("Upside / quote", f'IF({dref("px")}>0,B12/{dref("px")}-1,"n.a.")'),
        14: ("Terminal operating revenue", f"'Model'!{end}6"),
        15: ("Terminal credit ROE on closing equity", f"'Model'!{end}14/'Model'!{end}15"),
    }
    implied = (
        f"(1+{dref('cbgt')}/2)*({dref('nimat')}-{dref('cropex')})*(1-{dref('tax')})/{dref('cap')}"
    )
    valid = f"AND({dref('cap')}>0,{dref('cap')}<=1,{dref('tax')}>=0,{dref('tax')}<1,{dref('cgt')}>=0,{dref('cke')}>{dref('cgt')},{dref('croe')}>{dref('cgt')},ABS({dref('cbgt')}-{dref('cgt')})<=0.000000001,ABS(({implied})-{dref('croe')})<=MAX(0.000001,ABS({implied})*0.000001),{dref('sh')}>0,{dref('years')}={n},'Dashboard'!B44>=1,'Dashboard'!B44<=3,'Dashboard'!B44=INT('Dashboard'!B44))"
    put("Valuation", "A16", "Input consistency")
    put("Valuation", "B16", f'=IF({valid},"Valid","Invalid inputs; review assumptions")')
    for row, (label, formula) in valuation.items():
        put("Valuation", f"A{row}", label)
        put("Valuation", f"B{row}", "=" + formula, PCT if row in {13, 15} else NUM2)
    for key, row in (
        ("opev", 5),
        ("crev", 9),
        ("eqv", 11),
        ("vps", 12),
        ("up", 13),
        ("termrev", 14),
        ("croet", 15),
    ):
        put("Dashboard", f"A{R[key]}", valuation[row][0])
        put(
            "Dashboard",
            f"B{R[key]}",
            f"='Valuation'!B{row}",
            PCT if key in {"up", "croet"} else NUM2,
        )
    put("Dashboard", "D45", "Valuation shares (M)")
    for column, label in (("E", "Base"), ("F", "Bear"), ("G", "Bull")):
        put("Dashboard", f"{column}45", label)
    put("Dashboard", "D46", "='Valuation'!B12")
    for col, index in (("E", 1), ("F", 2), ("G", 3)):
        put("Dashboard", f"{col}46", index, "0")
    for row, multiplier in ((47, 0.95), (48, 1.0), (49, 1.05)):
        put("Dashboard", f"D{row}", f"=$B$32*{multiplier}", NUM2)
    put("Dashboard", "D53", "Discount sensitivity: WACC across; credit Ke down")
    put("Dashboard", "D54", "='Valuation'!B12")
    for col, delta in (("E", -0.01), ("F", 0.0), ("G", 0.01)):
        put("Dashboard", f"{col}54", f"=$B$19+({delta})", PCT)
    for row, delta in ((55, -0.01), (56, 0.0), (57, 0.01)):
        put("Dashboard", f"D{row}", f"=$B$27+({delta})", PCT)
    put("Scenario", "A1", "Live case comparison at the valuation-share assumption")
    put("Scenario", "A6", "Bear provenance")
    put("Scenario", "B6", BEAR_PROVENANCE_NOTES[bear_provenance])
    put("Scenario", "A2", "Case")
    put("Scenario", "E2", "Value / share ($)")
    for row, name, col in ((3, "Bear", "F"), (4, "Base", "E"), (5, "Bull", "G")):
        put("Scenario", f"A{row}", name)
        put("Scenario", f"E{row}", f"='Dashboard'!{col}48", NUM2)
    for row in (5, 9, 10, 11, 12, 14, 15):
        put("Scenario", f"A{row + 5}", valuation[row][0])
        put("Scenario", f"B{row + 5}", f"='Valuation'!B{row}", PCT if row == 15 else NUM2)
    expected = {
        f"Valuation!B{row}": getattr(m, attr)
        for row, attr in (
            (2, "pv_op_fcff"),
            (3, "op_terminal_ev"),
            (4, "pv_op_terminal"),
            (5, "operating_ev"),
            (6, "pv_credit_fcfe"),
            (7, "credit_terminal"),
            (8, "pv_credit_terminal"),
            (9, "credit_equity_value"),
            (11, "equity_value"),
            (12, "vps"),
            (14, "op_terminal_revenue"),
            (15, "credit_terminal_roe"),
        )
    }
    for col, name in (("E", "base"), ("F", "bear"), ("G", "bull")):
        output = mirror(cases[name])
        for row, multiplier in ((47, 0.95), (48, 1.0), (49, 1.05)):
            expected[f"Dashboard!{col}{row}"] = output.vps / multiplier
    for row, name in ((3, "bear"), (4, "base"), (5, "bull")):
        expected[f"Scenario!E{row}"] = mirror(cases[name]).vps
    for col, dw in (("E", -0.01), ("F", 0.0), ("G", 0.01)):
        for row, dk in ((55, -0.01), (56, 0.0), (57, 0.01)):
            expected[f"Dashboard!{col}{row}"] = mirror(
                replace(s, wacc=s.wacc + dw, credit_ke=s.credit_ke + dk)
            ).vps
    for year, row in enumerate(m.rows, 3):
        c = get_column_letter(year)
        for r, attr in (
            (4, "comm_rev"),
            (5, "fpay_rev"),
            (6, "op_rev"),
            (7, "op_ebit"),
            (8, "op_da"),
            (9, "op_capex"),
            (10, "op_fcff"),
            (11, "df_op"),
            (13, "cb"),
            (14, "credit_ni"),
            (15, "reqcap"),
            (16, "credit_fcfe"),
            (17, "df_cr"),
        ):
            expected[f"Model!{c}{r}"] = getattr(row, attr)
    tables = [
        {"sheet": "Dashboard", "range": "D46:G49", "row_input": "B44", "column_input": "B32"},
        {"sheet": "Dashboard", "range": "D54:G57", "row_input": "D19", "column_input": "D27"},
    ]
    if (
        input_receipt is not None
        and (input_receipt.request.recipe_context or {}).get("diagnostics")
        == "meli-report-driver-checks/v1"
    ):
        calculations = {row.key: row.value for row in input_receipt.calculations}
        if not set(MELI_DIAGNOSTIC_UNITS).issubset(calculations):
            raise InputEvidenceError("workbook_driver_diagnostics_missing")
        observed = {row.key: row.value for row in input_receipt.inputs}
        put("Scenario", "A24", "Base allocation check — USD millions; not reported segment results")
        put("Scenario", "A25", "Reported TTM operating income")
        put("Scenario", "B25", calculations["operating_income"], NUM2)
        diagnostics = (
            (
                26,
                "Commerce profit proxy",
                "='Actuals'!B3*'Assumptions'!B6",
                "commerce_profit_proxy_year0",
                NUM2,
            ),
            (
                27,
                "Non-credit payments profit proxy",
                "='Actuals'!B8*'Assumptions'!B11",
                "payments_profit_proxy_year0",
                NUM2,
            ),
            (
                29,
                "Credit profit proxy",
                "='Actuals'!B20*(B28-'Assumptions'!B25)",
                "credit_profit_proxy_year0",
                NUM2,
            ),
            (30, "Uncalibrated combined proxy", "=B26+B27+B29", "profit_proxy_year0", NUM2),
            (31, "Proxy less reported operating income", "=B30-B25", "profit_proxy_variance", NUM2),
            (
                32,
                "Proxy variance / reported operating income",
                "=B31/B25",
                "profit_proxy_variance_fraction",
                PCT,
            ),
            (
                33,
                "Residual credit profit calibration",
                "=B25-B26-B27",
                "residual_credit_profit_calibration",
                NUM2,
            ),
            (
                34,
                "Implied residual credit cost / ending book",
                "=B28-B33/'Actuals'!B20",
                "credit_cost_calibration",
                PCT,
            ),
            (35, "Sum after residual allocation", "=B26+B27+B33", "calibrated_profit_sum", NUM2),
            (
                36,
                "Residual allocation arithmetic difference",
                "=B35-B25",
                "calibrated_profit_variance",
                NUM2,
            ),
            (
                39,
                "Selected-case forward year-one profit proxy",
                "='Model'!C7+'Model'!C14/(1-'Dashboard'!B17)",
                "forward_profit_year1",
                NUM2,
            ),
            (
                40,
                "Forward year one less reported TTM profit",
                "=B39-B25",
                "forward_profit_variance",
                NUM2,
            ),
        )
        put("Scenario", "A28", "Reported annualized NIMAL proxy; retain its original period")
        put("Scenario", "B28", observed["nimal_actual"], PCT)
        for row, label, formula, key, fmt in diagnostics:
            put("Scenario", f"A{row}", label)
            put("Scenario", f"B{row}", formula, fmt)
            expected[f"Scenario!B{row}"] = calculations[key]
        put(
            "Scenario", "A37", "Arithmetic tolerance ($M); not a tolerance for economic assumptions"
        )
        put("Scenario", "B37", 0.000001)
        put(
            "Scenario",
            "A38",
            "Reported NIMAL period and ending book are proxy bases; residual costs are unreported.",
        )
        put(
            "Scenario",
            "A42",
            "Forward growth/margins/costs are assumptions; the bridge does not prove segment profitability.",
        )
        for index, (prefix, key, values) in enumerate(BRIDGE_SENSITIVITY_CASES):
            start = 45 + index * 9
            baseline = {
                "credit_funding_debt_allocation": s.credit_funding_debt_allocation,
                "financing_other_expenses_reserve": s.financing_other_expenses_reserve,
                "credit_cash_allocation": s.credit_cash_allocation,
                "operating_cash_reserve": s.operating_cash_reserve,
            }[key]
            put("Scenario", f"A{start}", f"{key}: alternative allocation ($M) / value per share")
            put("Scenario", f"A{start + 1}", "Qualified allocation baseline (fixed)")
            put("Scenario", f"B{start + 1}", baseline, NUM2)
            sign = 1 if key == "credit_funding_debt_allocation" else -1
            for offset, value in enumerate(values, start + 2):
                put("Scenario", f"A{offset}", value, NUM2)
                put(
                    "Scenario",
                    f"B{offset}",
                    f"='Valuation'!B12+({value}-B{start + 1})*{sign}/'Dashboard'!B32",
                    NUM2,
                )
                expected[f"Scenario!B{offset}"] = calculations[f"{prefix}_{value}_vps"]
        put(
            "Scenario",
            "A82",
            "Cash tables change one allocation; business value and the qualified source baseline stay fixed.",
        )
        put(
            "Scenario",
            "A83",
            "Other forecast edits flow through the selected case; source/allocation changes require requalification.",
        )
        # Couple the credit steady state when a user tests credit operating costs.
        for column in ("B", "C", "D"):
            put("Assumptions", f"{column}28", f"={column}22", PCT)
            put(
                "Assumptions",
                f"{column}29",
                f"=(1+{column}22/2)*({column}24-{column}25)*(1-{column}17)/{column}26",
                PCT,
            )
            for row in (28, 29):
                sheets["Assumptions"]["inputs"].remove(f"{column}{row}")
        put("Dashboard", "C25", "Credit cost override")
        put(
            "Dashboard",
            "B25",
            "=IF(D25=\"\",CHOOSE($B$44,'Assumptions'!B25,'Assumptions'!C25,'Assumptions'!D25),D25)",
            PCT,
        )
        put("Dashboard", "B28", "=B22", PCT)
        put("Dashboard", "B29", "=(1+B22/2)*(B24-B25)*(1-B17)/B26", PCT)
        sheets["Dashboard"]["inputs"].append("D25")
        sheets["Dashboard"]["formats"]["D25"] = PCT
        put("Dashboard", "D62", "Base-case credit operating cost stress; coupled terminal ROE")
        put("Dashboard", "D64", "='Valuation'!B12")
        put("Dashboard", "E64", 1, "0")
        costs = (
            ("600bp", 0.06),
            ("800bp", 0.08),
            ("1000bp", 0.10),
            ("1200bp", 0.12),
            ("calibrated", calculations["credit_cost_calibration"]),
        )
        for row, (_label, cost) in enumerate(costs, 65):
            put("Dashboard", f"D{row}", cost, PCT)
            numeric = asdict(s) | {"credit_opex_ratio": cost}
            numeric_inputs = {
                key: value for key, value in numeric.items() if isinstance(value, (float, int))
            }
            roe = implied_credit_terminal_roe(numeric_inputs)
            expected[f"Dashboard!E{row}"] = mirror(
                replace(s, credit_opex_ratio=cost, credit_terminal_roe=roe)
            ).vps
        tables.append(
            {"sheet": "Dashboard", "range": "D64:E69", "row_input": "B44", "column_input": "D25"}
        )
    return {
        "schema": "meli-workbook-spec/v1",
        "sheets": list(sheets.values()),
        "data_tables": tables,
        "expected": expected,
        "input_receipt_sha256": digest,
        "scenario_commitment_sha256": scenario_commitment,
        "input_recipe": recipe,
        "bear_provenance_note": BEAR_PROVENANCE_NOTES[bear_provenance],
    }


def build_workbook(
    s: Assum,
    m: Mirror,
    dest: Path,
    *,
    cases: dict[str, Assum],
    runtime: WorkbookRuntime | None = None,
    input_receipt: ModelInputReceipt | None = None,
    ticker: str = "MELI",
    bear_provenance: BearProvenance = "explicit_vector_acceptance_separate",
) -> None:
    """Author, recalculate and verify before replacing the requested workbook."""
    spec = workbook_spec(
        s,
        m,
        cases=cases,
        input_receipt=input_receipt,
        ticker=ticker,
        bear_provenance=bear_provenance,
    )
    runtime = runtime or WorkbookRuntime.from_environment()
    dest.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".meli-workbook-", dir=dest.parent) as temporary:
        work = Path(temporary)
        (work / "node_modules").symlink_to(runtime.node_modules.resolve(), target_is_directory=True)
        adapter = Path(__file__).resolve().parents[2] / "execution" / "meli_workbook_adapter.mjs"
        shutil.copyfile(adapter, work / "adapter.mjs")
        payload = json.dumps(spec, sort_keys=True, separators=(",", ":"), allow_nan=False)
        (work / "spec.json").write_text(payload, encoding="utf-8")
        try:
            result = subprocess.run(
                [
                    str(runtime.node),
                    str(work / "adapter.mjs"),
                    str(work / "spec.json"),
                    str(work / "workbook.xlsx"),
                ],
                capture_output=True,
                text=True,
                check=False,
                timeout=120,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise InputEvidenceError("workbook_artifact_engine_unavailable") from exc
        if result.returncode:
            raise InputEvidenceError("workbook_artifact_engine_failed")
        receipt_path = work / "workbook.xlsx.recalculation.json"
        if not receipt_path.is_file() or receipt_path.stat().st_size > 1024 * 1024:
            raise InputEvidenceError("workbook_recalculation_receipt_missing")
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        output = work / "workbook.xlsx"
        if (
            not isinstance(receipt, dict)
            or receipt.get("status") != "verified"
            or receipt.get("spec_sha256") != hashlib.sha256(payload.encode()).hexdigest()
            or receipt.get("workbook_sha256") != hashlib.sha256(output.read_bytes()).hexdigest()
        ):
            raise InputEvidenceError("workbook_recalculation_commitment_mismatch")
        os.replace(output, dest)
        os.replace(receipt_path, dest.with_suffix(dest.suffix + ".recalculation.json"))
