"""Real canonical writer/reader contract and exact DCF actual admission."""

from __future__ import annotations

import os
import shutil
import sqlite3
import subprocess
import sys
from collections.abc import Callable
from datetime import UTC, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest

from dcf import redesign
from dcf.primary_fact_overlay import Statement
from provenance.fact_plane_v2 import FactCellV2
from provenance.fact_read_model import FactReadModel, ProvenanceBundleRead
from provenance.metric_ontology import MetricOntology
from sources.dcf_statements import DcfStatementCell, DcfStatementInputs, read_dcf_statements
from tests.fixtures.dcf_statements import STAMP, seed_dcf_statements


@pytest.fixture(scope="module")
def statement_database(
    migrated_db: Callable[..., Path], tmp_path_factory: pytest.TempPathFactory
) -> Path:
    path = migrated_db(tmp_path_factory.mktemp("canonical_dcf") / "statements.sqlite")
    income: list[dict[str, object]] = []
    balance: list[dict[str, object]] = []
    cashflow: list[dict[str, object]] = []
    for quarter in range(1, 5):
        coordinate: dict[str, object] = {"fiscalYear": 2025, "period": f"Q{quarter}"}
        income.append(
            {
                **coordinate,
                "revenue": 1000,
                "costOfRevenue": 500,
                "researchAndDevelopmentExpenses": 100,
                "sellingGeneralAndAdministrativeExpenses": 100,
                "operatingIncome": 300,
                "netIncome": 250,
                "weightedAverageShsOutDil": 100,
            }
        )
        balance.append(
            {
                **coordinate,
                "totalStockholdersEquity": 1000,
                "cashAndCashEquivalents": 50,
                "totalDebt": 25,
                "financeLeaseLiability": 0,
            }
        )
        cashflow.append(
            {
                **coordinate,
                "depreciationAndAmortization": 25,
                "stockBasedCompensation": 10,
                "capitalExpenditure": 0,
            }
        )
    with sqlite3.connect(path) as conn:
        seed_dcf_statements(
            conn,
            "TEST",
            {"income": income, "balance": balance, "cash_flow": cashflow},
            currency="USD",
            publication_granularity="all",
        )
    return path


@pytest.fixture(scope="module")
def admitted(statement_database: Path) -> DcfStatementInputs:
    with sqlite3.connect(statement_database) as conn:
        return read_dcf_statements(conn, "TEST", as_of=STAMP)


def test_canonical_round_trip_preserves_reported_zero_and_exact_bridge_lineage(
    admitted: DcfStatementInputs,
) -> None:
    admitted.require_complete_actuals()
    assert all(item.available for item in admitted.cells)
    assert all(row["capitalExpenditure"] == 0 for row in admitted.builder_records("cash_flow"))
    lineage = admitted.bridge_lineage("balance")
    assert lineage["status"] == "ok"
    assert all(
        item.provenance and item.definition_revision_id and item.resolution_revision_id
        for item in admitted.cells
    )
    restored = DcfStatementInputs.model_validate_json(admitted.model_dump_json())
    assert restored == admitted


def test_statement_reader_uses_one_sealed_provenance_batch(
    statement_database: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    calls: list[tuple[str, ...]] = []
    original = FactReadModel.provenance_bundles

    def spy(
        reader: FactReadModel, observation_ids: tuple[str, ...], *, cutoff: datetime
    ) -> tuple[ProvenanceBundleRead, ...]:
        calls.append(observation_ids)
        return original(reader, observation_ids, cutoff=cutoff)

    monkeypatch.setattr(FactReadModel, "provenance_bundles", spy)
    with sqlite3.connect(statement_database) as conn:
        inputs = read_dcf_statements(conn, "TEST", as_of=STAMP)
    inputs.require_complete_actuals()
    assert len(calls) == 1
    assert len(calls[0]) == len(inputs.cells)

    def serial(
        reader: FactReadModel, observation_ids: tuple[str, ...], *, cutoff: datetime
    ) -> tuple[ProvenanceBundleRead, ...]:
        return tuple(
            ProvenanceBundleRead(
                observation_id=observation_id,
                bundle=reader.provenance_bundle(observation_id, cutoff=cutoff),
            )
            for observation_id in observation_ids
        )

    monkeypatch.setattr(FactReadModel, "provenance_bundles", serial)
    with sqlite3.connect(statement_database) as conn:
        serial_inputs = read_dcf_statements(conn, "TEST", as_of=STAMP)
    assert serial_inputs.model_dump_json() == inputs.model_dump_json()


@pytest.mark.parametrize(
    "concept",
    [
        "capital_expenditure",
        "cost_of_revenue",
        "research_and_development",
        "total_debt",
        "cash_and_equivalents",
        "weighted_avg_shares_diluted",
    ],
)
def test_missing_actual_never_becomes_zero(admitted: DcfStatementInputs, concept: str) -> None:
    incomplete = admitted.model_copy(
        update={"cells": tuple(item for item in admitted.cells if item.concept != concept)}
    )
    with pytest.raises(ValueError, match=r"missing|unavailable"):
        incomplete.require_complete_actuals()


def test_duplicate_semantically_unbound_cells_are_ambiguous(admitted: DcfStatementInputs) -> None:
    duplicate = admitted.model_copy(update={"cells": (*admitted.cells, admitted.cells[0])})
    with pytest.raises(ValueError, match="ambiguous"):
        duplicate.require_complete_actuals()


def test_missing_middle_quarter_cannot_be_summed_as_a_full_year(
    admitted: DcfStatementInputs,
) -> None:
    incomplete = admitted.model_copy(
        update={
            "cells": tuple(
                item
                for item in admitted.cells
                if item.provenance and item.provenance.cell.fiscal_period != "Q2"
            )
        }
    )
    with pytest.raises(ValueError, match="noncontiguous"):
        incomplete.require_complete_actuals()


def _change_coordinates(
    admitted: DcfStatementInputs, change: Callable[[DcfStatementCell], DcfStatementCell]
) -> DcfStatementInputs:
    return admitted.model_copy(update={"cells": tuple(map(change, admitted.cells))})


def test_off_calendar_fiscal_identity_is_retained(admitted: DcfStatementInputs) -> None:
    def shift(item: DcfStatementCell) -> DcfStatementCell:
        assert item.provenance
        source = item.provenance.cell
        # A 13-week fiscal calendar offset retains exact original fiscal labels.
        delta = timedelta(days=91)
        cell = source.model_copy(
            update={
                "period_start": source.period_start - delta if source.period_start else None,
                "period_end": source.period_end - delta,
            }
        )
        return item.model_copy(
            update={"provenance": item.provenance.model_copy(update={"cell": cell})}
        )

    shifted = _change_coordinates(admitted, shift)
    shifted.require_complete_actuals()
    row = shifted.builder_records("income")[0]
    assert row["fiscalYear"] == 2025 and row["period"] == "Q4"
    assert row["date"] == "2025-10-01"


def test_ytd_durations_cannot_masquerade_as_discrete_quarters(
    admitted: DcfStatementInputs,
) -> None:
    def change(item: DcfStatementCell) -> DcfStatementCell:
        assert item.provenance
        cell = item.provenance.cell
        if cell.period_kind != "duration":
            return item
        return item.model_copy(
            update={
                "provenance": item.provenance.model_copy(
                    update={
                        "cell": cell.model_copy(
                            update={"period_start": datetime(2025, 1, 1, tzinfo=UTC)}
                        )
                    }
                )
            }
        )

    with pytest.raises(ValueError, match=r"overlapping|unsupported fiscal duration"):
        _change_coordinates(admitted, change).require_complete_actuals()


def test_mixed_currency_actuals_are_not_added(admitted: DcfStatementInputs) -> None:
    def change(item: DcfStatementCell) -> DcfStatementCell:
        assert item.provenance
        if item.concept != "capital_expenditure":
            return item
        return item.model_copy(
            update={
                "provenance": item.provenance.model_copy(
                    update={
                        "observation": item.provenance.observation.model_copy(
                            update={"currency": "EUR"}
                        )
                    }
                )
            }
        )

    with pytest.raises(ValueError, match="currency"):
        _change_coordinates(admitted, change).require_complete_actuals()


def test_cutoff_excludes_later_publication_and_preserves_callers_transaction(
    statement_database: Path,
) -> None:
    with sqlite3.connect(statement_database) as conn:
        conn.execute("BEGIN")
        assert not read_dcf_statements(conn, "TEST", as_of=STAMP - timedelta(seconds=1)).cells
        assert conn.in_transaction
        assert conn.row_factory is None
    with (
        sqlite3.connect(statement_database) as conn,
        pytest.raises(ValueError, match="timezone-aware"),
    ):
        read_dcf_statements(conn, "TEST", as_of=datetime(2026, 9, 19))


@pytest.mark.parametrize("corruption", ["commitment", "publication_member"])
def test_tampered_immutable_source_is_unavailable(
    statement_database: Path, tmp_path: Path, corruption: str
) -> None:
    path = tmp_path / "tampered.sqlite"
    shutil.copyfile(statement_database, path)
    with sqlite3.connect(path) as conn:
        # Simulate storage corruption, not an allowed application write.
        if corruption == "commitment":
            conn.execute("DROP TRIGGER trg_fact_observation_payload_commitments_v2_append_only")
            conn.execute(
                "UPDATE fact_observation_payload_commitments_v2 SET observation_payload_sha256=? "
                "WHERE observation_id=(SELECT o.observation_id FROM fact_observations_v2 o "
                "JOIN fact_cells_v2 c ON c.fact_cell_id=o.fact_cell_id "
                "WHERE c.concept_name='revenue' LIMIT 1)",
                ("0" * 64,),
            )
        else:
            conn.execute("DROP TRIGGER trg_source_fact_publication_members_append_only")
            conn.execute(
                "UPDATE source_fact_publication_members SET canonical_member_sha256=? "
                "WHERE member_ordinal=0",
                ("0" * 64,),
            )
        corrupted = read_dcf_statements(conn, "TEST", as_of=STAMP)
    assert corrupted.cells
    assert any("canonical_evidence_invalid" in item.reason_codes for item in corrupted.cells)
    with pytest.raises(ValueError, match="unavailable"):
        corrupted.require_complete_actuals()


@pytest.mark.parametrize("case", ["scaled_unit", "non_primary", "future_period"])
def test_actual_resolver_rejects_inadmissible_coordinates(
    migrated_db: Callable[..., Path], tmp_path: Path, case: str
) -> None:
    path = migrated_db(tmp_path / "boundary.sqlite")

    def transform(cell: FactCellV2) -> FactCellV2:
        update: dict[str, object] = {"semantic_key_sha256": None}
        if case == "scaled_unit":
            update["unit_key"] = "USD millions"
        elif case == "future_period":
            update["period_start"] = datetime(2027, 1, 1, tzinfo=UTC)
            update["period_end"] = datetime(2027, 3, 31, tzinfo=UTC)
            update["fiscal_year"] = 2027
        return FactCellV2.model_validate({**cell.model_dump(), **update})

    with sqlite3.connect(path) as conn:
        seed_dcf_statements(
            conn,
            "TEST",
            {"income": [{"fiscalYear": 2025, "period": "Q1", "revenue": 1000}]},
            currency="USD",
            transform_cell=transform,
        )
        if case == "non_primary":
            conn.execute("UPDATE documents SET source_type='fmp', source_quality_tier='aggregator'")
        projection = read_dcf_statements(conn, "TEST", as_of=STAMP)
    assert len(projection.cells) == 1
    assert not projection.cells[0].available
    with pytest.raises(ValueError, match="unavailable"):
        projection.require_complete_actuals()


def test_definition_change_without_comparability_rejects_existing_actuals(
    statement_database: Path, tmp_path: Path
) -> None:
    path = tmp_path / "definition.sqlite"
    shutil.copyfile(statement_database, path)
    with sqlite3.connect(path) as conn:
        before = read_dcf_statements(conn, "TEST", as_of=STAMP)
        revenue = next(item for item in before.cells if item.concept == "revenue")
        ontology = MetricOntology(conn)
        definition = ontology.metric_definition_as_known(revenue.metric_id, STAMP)
        assert definition is not None
        later = STAMP + timedelta(days=1)
        ontology.persist_metric_definition(
            definition.model_copy(
                update={
                    "metric_definition_revision_id": "dcf-semantic-change",
                    "idempotency_key": "dcf-semantic-change",
                    "revision": definition.revision + 1,
                    "supersedes_metric_definition_revision_id": definition.metric_definition_revision_id,
                    "definition_text": "Revenue under a changed recognition perimeter",
                    "effective_at": later,
                    "knowledge_at": later,
                    "recorded_at": later,
                }
            )
        )
        after = read_dcf_statements(conn, "TEST", as_of=later)
        old = read_dcf_statements(conn, "TEST", as_of=STAMP)
    old.require_complete_actuals()
    assert any(
        "active_metric_definition_or_binding_unavailable" in item.reason_codes
        for item in after.cells
        if item.concept == "revenue"
    )
    with pytest.raises(ValueError, match="unavailable"):
        after.require_complete_actuals()


def test_actual_builder_preserves_zero_capex_and_emits_no_unavailable_gross_margin(
    statement_database: Path, tmp_path: Path
) -> None:
    repo = tmp_path / "repo"
    (repo / "data").mkdir(parents=True)
    database = repo / "data" / "canonical.sqlite"
    shutil.copyfile(statement_database, database)
    profile = repo / "data" / "historical" / "fmp" / "TEST_profile.json"
    profile.parent.mkdir(parents=True)
    profile.write_text('{"companyName":"Synthetic","price":50,"beta":1,"currency":"USD"}')
    (profile.parent / "TEST_geo_segments_annual.json").write_text(
        '[{"fiscalYear":2025,"period":"FY","data":{"United States":100}}]'
    )
    destination = repo / "synthetic.xlsx"
    completed = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).resolve().parents[1] / "execution" / "build_redesigned_dcf.py"),
        ],
        env={
            **os.environ,
            "DCF_TICKER": "TEST",
            "DCF_REPO_ROOT": str(repo),
            "DCF_DEST": str(destination),
            "EARNINGS_SUMMARY_DB_PATH": str(database),
        },
        capture_output=True,
        text=True,
        timeout=90,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    inputs = redesign.read_inputs(destination)
    assert inputs is not None and inputs.capex_2026_m == 0
    wb = openpyxl.load_workbook(destination)
    try:
        sheet = wb["Financials"]
        row = next(
            row
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row, 1).value == "    gross margin"
        )
        assert all(sheet.cell(row, column).value is None for column in range(2, 6))
    finally:
        wb.close()


@pytest.mark.parametrize(
    ("days", "available"), [(284, False), (420, False), (365, True), (371, True)]
)
def test_complete_fiscal_year_must_match_existing_annual_span(
    admitted: DcfStatementInputs, days: int, available: bool
) -> None:
    def resize(item: DcfStatementCell) -> DcfStatementCell:
        assert item.provenance
        cell = item.provenance.cell
        assert cell.fiscal_period is not None
        quarter = int(cell.fiscal_period[1])
        start = datetime(2025, 1, 1, tzinfo=UTC)
        boundaries = [days * position // 4 for position in range(5)]
        resized = cell.model_copy(
            update={
                "period_start": start + timedelta(days=boundaries[quarter - 1])
                if cell.period_kind == "duration"
                else None,
                "period_end": start + timedelta(days=boundaries[quarter] - 1),
            }
        )
        return item.model_copy(
            update={"provenance": item.provenance.model_copy(update={"cell": resized})}
        )

    projection = _change_coordinates(admitted, resize)
    if available:
        projection.require_complete_actuals()
    else:
        with pytest.raises(ValueError, match="annual span"):
            projection.require_complete_actuals()


@pytest.mark.parametrize("case", ["definition", "extended_partial_quarter"])
def test_actual_generic_builder_rejects_unbound_series_and_long_partial_quarter(
    admitted: DcfStatementInputs, migrated_db: Callable[..., Path], tmp_path: Path, case: str
) -> None:
    repo = tmp_path / case
    profile = repo / "data" / "historical" / "fmp" / "TEST_profile.json"
    profile.parent.mkdir(parents=True)
    profile.write_text('{"companyName":"Synthetic","price":50,"beta":1,"currency":"USD"}')
    database = migrated_db(repo / "data" / "canonical.sqlite")
    rows: dict[Statement, list[dict[str, object]]] = {
        statement: admitted.builder_records(statement)
        for statement in ("income", "balance", "cash_flow")
    }
    if case == "extended_partial_quarter":
        for values in rows.values():
            values.append({**values[-1], "fiscalYear": 2026, "period": "Q1"})

    def transform(cell: FactCellV2) -> FactCellV2:
        update: dict[str, object] = {"semantic_key_sha256": None}
        if (
            case == "definition"
            and cell.concept_name == "revenue"
            and cell.fiscal_period in {"Q3", "Q4"}
        ):
            update["taxonomy_name"] = "changed-reported-financial-definition"
        if case == "extended_partial_quarter" and cell.fiscal_year == 2026:
            update["period_end"] = datetime(2026, 6, 30, tzinfo=UTC)
        return FactCellV2.model_validate({**cell.model_dump(), **update})

    with sqlite3.connect(database) as conn:
        seed_dcf_statements(conn, "TEST", rows, currency="USD", transform_cell=transform)
    destination = repo / "rejected.xlsx"
    completed = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).resolve().parents[1] / "execution" / "build_redesigned_dcf.py"),
        ],
        env={
            **os.environ,
            "DCF_TICKER": "TEST",
            "DCF_REPO_ROOT": str(repo),
            "DCF_DEST": str(destination),
            "EARNINGS_SUMMARY_DB_PATH": str(database),
        },
        capture_output=True,
        text=True,
        timeout=90,
        check=False,
    )
    assert completed.returncode != 0
    assert (
        "continuity unavailable" if case == "definition" else "unsupported fiscal duration"
    ) in completed.stderr
    assert not destination.exists()
    assert not (repo / "data" / "dcf_assumptions").exists()


@pytest.mark.parametrize("long_partial", [False, True])
def test_actual_discrete_half_year_history_and_malformed_partial_half(
    admitted: DcfStatementInputs,
    migrated_db: Callable[..., Path],
    tmp_path: Path,
    long_partial: bool,
) -> None:
    rows: dict[Statement, list[dict[str, object]]] = {
        statement: admitted.builder_records(statement)
        for statement in ("income", "balance", "cash_flow")
    }
    for records in rows.values():
        for record, (year, period) in zip(
            records, [(2024, "Q2"), (2024, "Q4"), (2025, "Q2"), (2025, "Q4")], strict=True
        ):
            record.update(fiscalYear=year, period=period)
        if long_partial:
            records.append({**records[-1], "fiscalYear": 2026, "period": "Q2"})

    def transform(cell: FactCellV2) -> FactCellV2:
        if cell.fiscal_year != 2026:
            return cell
        return FactCellV2.model_validate(
            {
                **cell.model_dump(),
                "semantic_key_sha256": None,
                "period_end": datetime(2026, 7, 20, tzinfo=UTC),
            }
        )

    database = migrated_db(tmp_path / "half-year.sqlite")
    with sqlite3.connect(database) as conn:
        seed_dcf_statements(conn, "TEST", rows, currency="USD", transform_cell=transform)
        projection = read_dcf_statements(conn, "TEST", as_of=STAMP)
    if long_partial:
        with pytest.raises(ValueError, match="unsupported fiscal duration"):
            projection.require_complete_actuals()
    else:
        projection.require_complete_actuals()
        assert len(projection.builder_records("income")) == 4
