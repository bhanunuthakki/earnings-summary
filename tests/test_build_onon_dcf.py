"""The builder requires an explicit verified artifact and stages its workbook."""

import json
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import pytest

from dcf.artifact_promotion import StagedFilePromotion
from dcf.input_evidence import InputEvidenceError
from dcf.onon_model import replay
from execution import build_onon_dcf as builder
from tests.test_onon_inputs import proof
from tests.test_onon_model import memo_inputs


def test_explicit_authority_and_dispatch_hash_required(tmp_path: Path) -> None:
    with pytest.raises(InputEvidenceError, match="missing_or_invalid"):
        builder.read_request(tmp_path / "missing.json")
    artifact = tmp_path / "inputs.json"
    artifact.write_text(json.dumps({"input_evidence": proof().request.model_dump(mode="json")}))
    with pytest.raises(InputEvidenceError, match="changed_after_dispatch"):
        builder.read_request(artifact, expected_sha256="f" * 64)
    request, _, digest = builder.read_request(artifact)
    assert request.ticker == "ONON"
    assert len(digest) == 64


@pytest.mark.parametrize("quote", [30.85, 0.1, 500.0])
def test_workbook_contains_full_replay_and_receipt(tmp_path: Path, quote: float) -> None:
    inputs = memo_inputs()
    inputs["price_usd"] = quote
    output = replay(inputs)
    destination = tmp_path / "stage.xlsx"
    builder.build_workbook(inputs, output, proof(), destination)
    wb = openpyxl.load_workbook(destination, data_only=True)
    assert set(wb.sheetnames) == {
        "Summary",
        "Inputs",
        "Reported inputs",
        "bear",
        "base",
        "bull",
        "Sensitivity",
        "Stresses",
        "Reverse DCF",
        "Evidence",
    }
    assert wb["Summary"]["B2"].value == pytest.approx(output["vps"])
    assert wb["base"]["J2"].value == pytest.approx(
        output["scenarios"]["base"]["rows"][0]["economic_fcff_chf_m"]
    )
    assert wb["Evidence"]["B2"].value == output["effective_inputs_sha256"]
    reverse: dict[str, object] = {}
    for row in wb["Reverse DCF"].iter_rows(min_row=2, values_only=True):
        assert len(row) == 2
        assert isinstance(row[0], str)
        reverse[row[0]] = row[1]
    assert reverse.keys() == output["reverse_dcf"].keys()
    for key, expected in output["reverse_dcf"].items():
        assert reverse[key] == (
            pytest.approx(expected, rel=1e-14, abs=1e-14)
            if isinstance(expected, float)
            else expected
        )
    if quote in {0.1, 500.0}:
        assert reverse["required_constant_2027_2031_growth"] is None
        assert reverse["growth_status"] in {"target_below_bracket", "target_above_bracket"}
    wb.close()


def test_snapshot_never_synthesizes_scenario_review() -> None:
    inputs = memo_inputs()
    result = builder.snapshot_payload(
        inputs,
        builder.model_output(inputs),
        proof(),
        Path("/tmp/ONON.xlsx"),
        Path("/tmp/inputs.json"),
        "a" * 64,
    )
    assert result["model"] == "onon_economic_fcff"
    assert result["scenarios"] == builder.model_output(inputs)["scenarios"]
    assert "scenario_acceptance" not in result


def test_market_observation_must_be_explicit_and_aware() -> None:
    with pytest.raises(InputEvidenceError, match="market_observation"):
        builder.market_clock({})
    with pytest.raises(InputEvidenceError, match="market_observation"):
        builder.market_clock({"market_observed_at": "2026-10-02T20:00:00"})
    assert builder.market_clock({"market_observed_at": "2026-10-02T20:00:00Z"}) == datetime(
        2026, 10, 2, 20, tzinfo=UTC
    )


def test_explicit_prior_receipt_preserves_review_clock(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    import hashlib
    import sqlite3
    from collections.abc import Generator
    from contextlib import contextmanager

    from dcf.input_evidence import ModelInputReceipt
    from tests.test_onon_inputs import NOW, bridge_proof

    values, receipt = bridge_proof()
    assumptions = dict(receipt.request.assumptions)
    assumptions["price_usd"] = assumptions["price_usd"].model_copy(
        update={"source_as_of": datetime(2026, 10, 2, tzinfo=UTC).date()}
    )
    request = receipt.request.model_copy(update={"assumptions": assumptions})
    artifact = tmp_path / "assumptions.json"
    artifact.write_text(
        json.dumps(
            {
                "input_evidence": request.model_dump(mode="json"),
                "market_observed_at": "2026-10-02T20:00:00Z",
            }
        )
    )
    receipt = receipt.model_copy(
        update={
            "request": request,
            "assumptions_source_path": str(artifact.resolve()),
            "assumptions_source_sha256": hashlib.sha256(artifact.read_bytes()).hexdigest(),
        }
    )
    verified: list[ModelInputReceipt] = []

    @contextmanager
    def connect(*args: object, **kwargs: object) -> Generator[sqlite3.Connection]:
        assert kwargs["role"] == builder.SQLiteConnectionRole.READ_ONLY
        with sqlite3.connect(":memory:") as conn:
            yield conn

    def prepared(*args: object, **kwargs: object) -> tuple[dict[str, float], ModelInputReceipt]:
        return values, receipt.model_copy(update={"verified_at": datetime(2026, 10, 4, tzinfo=UTC)})

    def verify(
        conn: sqlite3.Connection, supplied: ModelInputReceipt, **kwargs: object
    ) -> ModelInputReceipt:
        verified.append(supplied)
        return supplied

    monkeypatch.setattr(builder, "connect_sqlite", connect)

    def required(path: Path) -> Path:
        return path

    monkeypatch.setattr(builder, "require_db_path", required)
    monkeypatch.setattr(builder, "prepare_onon_inputs", prepared)
    monkeypatch.setattr(builder, "verify_onon_inputs", verify)
    _, reused, _ = builder.load_verified_inputs(
        db_path=tmp_path / "synthetic.db",
        assumptions_path=artifact,
        as_of=NOW,
        input_receipt=receipt,
    )
    assert reused == receipt
    assert verified == [receipt]
    with pytest.raises(InputEvidenceError, match="receipt_authority_mismatch"):
        builder.load_verified_inputs(
            db_path=tmp_path / "synthetic.db",
            assumptions_path=artifact,
            as_of=NOW,
            input_receipt=receipt.model_copy(update={"assumptions_source_sha256": "f" * 64}),
        )


def test_persist_rejects_source_mutation_before_database_access(tmp_path: Path) -> None:
    inputs = memo_inputs()
    with pytest.raises(InputEvidenceError, match="changed_after_review"):
        builder.persist_dcf_run(
            inputs,
            builder.model_output(inputs),
            proof(),
            db_path=tmp_path / "absent.db",
            assumptions_path=tmp_path / "inputs.json",
            destination=tmp_path / "stage.xlsx",
            repo_root=tmp_path,
            market_observed_at=datetime(2026, 10, 2, 20, tzinfo=UTC),
            calculated_at=datetime.now(UTC),
            artifact_promotion=StagedFilePromotion(tmp_path / "stage.xlsx", tmp_path / "live.xlsx"),
        )
