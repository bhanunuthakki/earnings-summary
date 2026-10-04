"""Verified ONON cash-rent/SBC FCFF workbook and staged persistence route.

Requires DCF_ONON_ASSUMPTIONS_PATH and configured database. The explicit artifact
owns input_evidence and a timezone-aware market_observed_at. Forecast sources
remain assumptions. DCF_ONON_SCENARIO_ACCEPTANCE_PATH optionally supplies an
independent analyst review; this builder never creates that review.

Run through execution/sqlite_bootstrap.py; the runtime owns import paths.
"""

from __future__ import annotations

import hashlib
import json
import os
from collections.abc import Mapping
from datetime import UTC, date, datetime
from pathlib import Path
from typing import cast

import openpyxl

from db_paths import configured_db_path, require_db_path
from dcf.artifact_promotion import (
    ArtifactPromotion,
    live_path_from_env,
    promotion_from_env,
    run_dcf_entrypoint,
)
from dcf.input_evidence import (
    InputEvidenceError,
    ModelInputReceipt,
    ModelInputRequest,
    canonical_digest,
)
from dcf.onon_inputs import (
    RECIPE,
    build_onon_equity_bridge,
    model_output,
    prepare_onon_inputs,
    verify_onon_inputs,
)
from dcf.onon_model import ASSUMPTION_KEYS
from dcf.persist import DcfRunRow, upsert
from dcf.provenance import build_file_provenance, schema_supports_provenance
from dcf.scenario_acceptance import ScenarioAcceptance, verify_scenario_acceptance
from sqlite_runtime import SQLiteConnectionRole, connect_sqlite

MODEL = "onon_economic_fcff"
ENGINE_VERSION = MODEL + "_v1"
REPO = Path(os.environ.get("DCF_REPO_ROOT") or Path(__file__).resolve().parents[1])
DEST = Path(os.environ.get("DCF_DEST") or REPO / "dcf" / "ONON.xlsx")


def read_request(
    path: Path, *, expected_sha256: str | None = None
) -> tuple[ModelInputRequest, dict[str, object], str]:
    try:
        raw = path.read_bytes()
        digest = hashlib.sha256(raw).hexdigest()
        if expected_sha256 is not None and expected_sha256 != digest:
            raise InputEvidenceError("assumptions_authority_changed_after_dispatch")
        payload = json.loads(raw)
        if not isinstance(payload, dict):
            raise ValueError("object_required")
        payload = cast(dict[str, object], payload)
        request = ModelInputRequest.model_validate(payload.get("input_evidence"))
    except (OSError, ValueError) as exc:
        if isinstance(exc, InputEvidenceError):
            raise
        raise InputEvidenceError("model_input_request_missing_or_invalid") from exc
    if request.recipe != RECIPE or request.ticker != "ONON":
        raise InputEvidenceError("onon_input_recipe_mismatch")
    return request, payload, digest


def market_clock(payload: Mapping[str, object]) -> datetime:
    raw = payload.get("market_observed_at")
    try:
        if not isinstance(raw, str):
            raise ValueError("explicit_clock_required")
        clock = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        if clock.tzinfo is None or clock.utcoffset() is None:
            raise ValueError("timezone_required")
    except ValueError as exc:
        raise InputEvidenceError("onon_market_observation_clock_required") from exc
    return clock


def load_verified_inputs(
    *,
    db_path: Path,
    assumptions_path: Path,
    as_of: datetime,
    expected_sha256: str | None = None,
    input_receipt: ModelInputReceipt | None = None,
) -> tuple[dict[str, float], ModelInputReceipt, dict[str, object]]:
    request, payload, digest = read_request(assumptions_path, expected_sha256=expected_sha256)
    if frozenset(request.assumptions) != ASSUMPTION_KEYS:
        raise InputEvidenceError("assumption_population_mismatch")
    values = {key: item.value for key, item in request.assumptions.items()}
    clock = market_clock(payload)
    valuation_date = date.fromordinal(int(values["valuation_date_ordinal"]))
    price_basis = request.assumptions["price_usd"]
    if clock > as_of or clock.date() > valuation_date or clock.date() != price_basis.source_as_of:
        raise InputEvidenceError("onon_market_observation_clock_mismatch")
    with connect_sqlite(require_db_path(db_path), role=SQLiteConnectionRole.READ_ONLY) as conn:
        conn.execute("BEGIN")
        values, receipt = prepare_onon_inputs(conn, request, effective_inputs=values, as_of=as_of)
        if input_receipt is not None:
            if (
                input_receipt.request != request
                or input_receipt.assumptions_source_path != str(assumptions_path.resolve())
                or input_receipt.assumptions_source_sha256 != digest
            ):
                raise InputEvidenceError("model_input_receipt_authority_mismatch")
            verify_onon_inputs(conn, input_receipt, effective_inputs=values, as_of=as_of)
            receipt = input_receipt
    return (
        values,
        receipt.model_copy(
            update={
                "assumptions_source_path": str(assumptions_path.resolve()),
                "assumptions_source_sha256": digest,
            }
        ),
        payload,
    )


def build_workbook(
    inputs: Mapping[str, float],
    output: Mapping[str, object],
    receipt: ModelInputReceipt,
    destination: Path,
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Summary"
    ws.append(["ONON economic FCFF", "USD per share / USD millions"])
    for key in ("vps", "equity_value", "operating_ev"):
        ws.append([key, output[key]])
    ws.append(["Model", MODEL])
    ws.append(
        [
            "Liability deduction",
            "Reported other nonlease financial liabilities; broader than borrowing debt",
        ]
    )
    ws.append(["Policy", "Cash rent and SBC expensed; leases excluded from debt bridge"])
    sheet = wb.create_sheet("Inputs")
    sheet.append(["Effective input", "Value"])
    for key, value in sorted(inputs.items()):
        sheet.append([key, value])
    sheet = wb.create_sheet("Reported inputs")
    sheet.append(
        [
            "Key",
            "CHF millions / share count millions",
            "Period start",
            "Period end",
            "Observation",
        ]
    )
    for item in receipt.inputs:
        sheet.append(
            [
                item.key,
                item.value,
                str(item.period_start or ""),
                item.period_end.isoformat(),
                item.reference.observation_id,
            ]
        )
    scenarios = cast(dict[str, dict[str, object]], output["scenarios"])
    fields = (
        "year",
        "revenue_chf_m",
        "growth",
        "adjusted_ebitda_margin",
        "post_rent_sbc_ebit_chf_m",
        "nopat_chf_m",
        "nonlease_da_chf_m",
        "capex_chf_m",
        "incremental_nwc_chf_m",
        "economic_fcff_chf_m",
    )
    for name, scenario in scenarios.items():
        sheet = wb.create_sheet(name)
        sheet.append(list(fields))
        for row in cast(list[dict[str, float]], scenario["rows"]):
            sheet.append([row[key] for key in fields])
        sheet.append([])
        for key in (
            "vps",
            "operating_ev",
            "equity_value",
            "gordon_terminal_weight",
            "exit_multiple_fair_value_usd",
        ):
            sheet.append([key, scenario[key]])
    sheet = wb.create_sheet("Sensitivity")
    sheet.append(["WACC", "Terminal growth", "USD per share"])
    for rate, cells in cast(dict[str, dict[str, float]], output["gordon_sensitivity"]).items():
        for growth, value in cells.items():
            sheet.append([float(rate), float(growth), value])
    for title, key in (("Stresses", "one_factor_stresses"), ("Reverse DCF", "reverse_dcf")):
        sheet = wb.create_sheet(title)
        sheet.append(["Driver", "Value"])
        for driver, value in cast(dict[str, float | str | None], output[key]).items():
            sheet.append([driver, value])
    sheet = wb.create_sheet("Evidence")
    sheet.append(["Commitment", "SHA256"])
    sheet.append(["effective_inputs", output["effective_inputs_sha256"]])
    sheet.append(["model_output", canonical_digest(dict(output))])
    sheet.append(["model_input_receipt", canonical_digest(receipt.model_dump(mode="json"))])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 44
        sheet.column_dimensions["B"].width = 42
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    wb.close()


def snapshot_payload(
    inputs: Mapping[str, float],
    output: Mapping[str, object],
    receipt: ModelInputReceipt,
    workbook: Path,
    assumptions_path: Path,
    digest: str,
) -> dict[str, object]:
    return {
        "model": MODEL,
        "effective_model_inputs": dict(inputs),
        "model_output": dict(output),
        "scenarios": output["scenarios"],
        "financial_period_end": receipt.request.financial_period_end.isoformat(),
        "value_per_share": output["vps"],
        "equity_value_m": output["equity_value"],
        "operating_ev_m": output["operating_ev"],
        "workbook": str(workbook),
        "assumption_provenance": {
            "authority": str(assumptions_path.resolve()),
            "sha256": digest,
            "rates": "dated_reviewed_effective_vector",
            "sync_status": "not_applicable",
        },
    }


def persist_dcf_run(
    inputs: Mapping[str, float],
    output: Mapping[str, object],
    receipt: ModelInputReceipt,
    *,
    db_path: Path,
    assumptions_path: Path,
    destination: Path,
    repo_root: Path,
    market_observed_at: datetime,
    calculated_at: datetime,
    artifact_promotion: ArtifactPromotion,
    scenario_acceptance: ScenarioAcceptance | None = None,
    scenario_acceptance_path: Path | None = None,
) -> bool:
    if (
        receipt.assumptions_source_path != str(assumptions_path.resolve())
        or hashlib.sha256(assumptions_path.read_bytes()).hexdigest()
        != receipt.assumptions_source_sha256
    ):
        raise InputEvidenceError("assumptions_authority_changed_after_review")
    if dict(output) != model_output(inputs):
        raise InputEvidenceError("model_output_input_mismatch")
    live = live_path_from_env(destination)
    snapshot = snapshot_payload(
        inputs, output, receipt, live, assumptions_path, receipt.assumptions_source_sha256 or ""
    )
    if scenario_acceptance is not None:
        verify_scenario_acceptance(
            scenario_acceptance,
            input_receipt=receipt,
            model=MODEL,
            engine_version=ENGINE_VERSION,
            effective_inputs=inputs,
            output=output,
            as_of=calculated_at,
            calculated_at=calculated_at,
        )
        if (
            scenario_acceptance_path is None
            or ScenarioAcceptance.model_validate_json(scenario_acceptance_path.read_bytes())
            != scenario_acceptance
        ):
            raise InputEvidenceError("scenario_acceptance_authority_changed")
    with connect_sqlite(
        require_db_path(db_path), role=SQLiteConnectionRole.WRITER, schema_preflight=True
    ) as conn:
        conn.execute("BEGIN IMMEDIATE")
        bridge = build_onon_equity_bridge(
            conn, receipt, effective_inputs=inputs, as_of=calculated_at
        )
        provenance = build_file_provenance(
            ticker="ONON",
            repo_root=repo_root,
            workbook_path=destination,
            workbook_locator_path=live,
            engine_version=ENGINE_VERSION,
            effective_inputs=dict(inputs),
            assumption_snapshot=snapshot,
            live_price=inputs["price_usd"],
            live_price_at=market_observed_at,
            live_price_source=receipt.request.assumptions["price_usd"].source_reference,
            source_files=((assumptions_path, "analyst_assumptions"),)
            + (
                ((scenario_acceptance_path, "analyst_scenario_acceptance"),)
                if scenario_acceptance_path
                else ()
            ),
            model_input_receipt=receipt.model_dump(mode="json"),
            equity_bridge_receipt=bridge.to_dict(),
            scenario_acceptance=scenario_acceptance.model_dump(mode="json")
            if scenario_acceptance is not None
            else None,
        )
        row = DcfRunRow(
            ticker="ONON",
            valuation_date=date.fromordinal(int(inputs["valuation_date_ordinal"])),
            horizon_years=5,
            wacc=inputs["wacc"],
            npv=float(cast(float, output["equity_value"])),
            npv_per_share=float(cast(float, output["vps"])),
            shares_outstanding=inputs["shares"] * 1e6,
            currency="USD",
            live_price=inputs["price_usd"],
            live_price_at=market_observed_at,
            mos_bar_used=None,
            assumption_snapshot_json=json.dumps(snapshot, allow_nan=False),
            notes="ONON cash-rent/SBC economic FCFF",
            provenance=provenance,
            calculated_at=calculated_at,
        )
        if not schema_supports_provenance(conn):
            raise InputEvidenceError("model_input_receipt_schema_unavailable")
        if (
            hashlib.sha256(assumptions_path.read_bytes()).hexdigest()
            != receipt.assumptions_source_sha256
        ):
            raise InputEvidenceError("assumptions_authority_changed_after_review")
        return upsert(conn, row, artifact_promotion=artifact_promotion)


def _main_owned() -> int:
    if os.environ.get("DCF_PERSIST", "1") != "1":
        raise InputEvidenceError("onon_verified_persistence_route_required")
    if os.environ.get("DCF_TICKER", "ONON") != "ONON":
        raise InputEvidenceError("onon_ticker_required")
    promotion = promotion_from_env(DEST)
    if promotion is None or DEST.resolve() == live_path_from_env(DEST).resolve():
        raise InputEvidenceError("atomic_artifact_promotion_required_use_refresh_dcf")
    raw = os.environ.get("DCF_ONON_ASSUMPTIONS_PATH", "").strip()
    if not raw:
        raise InputEvidenceError("explicit_assumptions_authority_required")
    assumptions_path = Path(raw)
    db = require_db_path(configured_db_path(REPO))
    prior_path = os.environ.get("DCF_ONON_MODEL_INPUT_RECEIPT_PATH", "").strip()
    prior_receipt = (
        ModelInputReceipt.model_validate_json(Path(prior_path).read_bytes()) if prior_path else None
    )
    inputs, receipt, payload = load_verified_inputs(
        db_path=db,
        assumptions_path=assumptions_path,
        as_of=datetime.now(UTC),
        expected_sha256=os.environ.get("DCF_ONON_ASSUMPTIONS_SHA256"),
        input_receipt=prior_receipt,
    )
    output = model_output(inputs)
    reviewed = os.environ.get("DCF_ONON_SCENARIO_ACCEPTANCE_PATH", "").strip()
    review_path = Path(reviewed) if reviewed else None
    review = (
        ScenarioAcceptance.model_validate_json(review_path.read_bytes()) if review_path else None
    )
    build_workbook(inputs, output, receipt, DEST)
    persisted = persist_dcf_run(
        inputs,
        output,
        receipt,
        db_path=db,
        assumptions_path=assumptions_path,
        destination=DEST,
        repo_root=REPO,
        market_observed_at=market_clock(payload),
        calculated_at=datetime.now(UTC),
        artifact_promotion=promotion,
        scenario_acceptance=review,
        scenario_acceptance_path=review_path,
    )
    print(
        f"RESULT\tONON\tvalue/sh=${float(cast(float, output['vps'])):.2f}\tdcf_runs={'ok' if persisted else 'skip'}"
    )
    return 0


def main() -> int:
    return run_dcf_entrypoint(
        REPO, "ONON", _main_owned, owner="build-onon-dcf", require_database=True
    )


if __name__ == "__main__":
    raise SystemExit(main())
